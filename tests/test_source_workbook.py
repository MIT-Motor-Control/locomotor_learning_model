"""Source-data workbook: cell values, layout and provenance checks."""

import hashlib
import importlib.util
import json
import sys
from pathlib import Path

import pytest
from openpyxl import load_workbook

SCRIPTS = Path(__file__).resolve().parents[1] / "scripts"
sys.path.insert(0, str(SCRIPTS))
spec = importlib.util.spec_from_file_location(
    "source_workbook", SCRIPTS / "build_source_workbook.py"
)
builder = importlib.util.module_from_spec(spec)
spec.loader.exec_module(builder)
sys.path.pop(0)


@pytest.fixture
def reproduction(tmp_path):
    """Return a reproduction folder with three Figure 2C tables and a report of hashes and statuses."""
    root = tmp_path / "reproduction"
    tables = root / "tables"
    tables.mkdir(parents=True)
    contents = {
        "figure2c_energy_minimization.csv": "delta_mps,tied_energy,split_energy,energy_change_pct\n0.2,3.147,3.161,0.4448681283762388\n",
        "figure2c_experimental_participants.csv": "condition,participant,energy_change_pct,note,missing,identifier,flag\n1,1,-8.12345678901234,=1+1,nan,12345678901234567,True\n",
        "figure2c_experimental_summary.csv": "delta_mps,mean_energy_change_pct\n0.4,-8.12345678901234\n",
    }
    artifacts = []
    for name, content in contents.items():
        (tables / name).write_text(content)
        artifacts.append(
            {"path": f"tables/{name}", "sha256": hashlib.sha256(content.encode()).hexdigest()}
        )
    report = {
        "schema_version": 1,
        "created_at": "2026-10-15",
        "inputs": {"manifest_sha256": "test-input-hash", "errors": []},
        "artifacts": artifacts,
        "results": [
            {"id": name, "status": "partial" if name == "figure2" else "blocked"}
            for name in ("figure1", "figure2", "figure3", "s1")
        ],
    }
    (root / "report.json").write_text(json.dumps(report))
    return root


def test_workbook_roundtrip_preserves_values_literals_and_partial_coverage(reproduction, tmp_path):
    output = tmp_path / "sources.xlsx"
    result = builder.build(reproduction, output)
    assert result["status"] == "passed"
    assert result["tables"] == 3
    workbook = load_workbook(output)
    metadata = json.loads(output.with_suffix(".validation.json").read_text())
    table = next(t for t in metadata["tables"] if t["csv"] == "figure2c_experimental_participants.csv")
    sheet = workbook[table["sheet"]]
    row = table["firstDataRow"]
    assert sheet.cell(row, 4).value == "=1+1"
    assert sheet.cell(row, 4).data_type == "s"
    assert sheet.cell(row, 5).value == "nan"
    assert sheet.cell(row, 6).value == "12345678901234567"
    assert sheet.cell(row, 7).value is True
    assert any(r["values"][:2] == ["Figure 2C model", "unavailable"] for r in metadata["guide"])
    assert {"Energy minimization", "Fig2C experiment"} <= set(workbook.sheetnames)
    assert "Fig2C model trials" not in workbook.sheetnames
    workbook.close()
    with pytest.raises(FileExistsError):
        builder.build(reproduction, output)


def test_workbook_rejects_csv_changed_since_report(reproduction, tmp_path):
    path = reproduction / "tables/figure2c_experimental_participants.csv"
    path.write_text(path.read_text().replace("-8.123", "-9.123"))
    with pytest.raises(ValueError, match="does not match"):
        builder.build(reproduction, tmp_path / "sources.xlsx")
    assert not (tmp_path / "sources.xlsx").exists()


def test_passed_figure_requires_its_model_tables(reproduction, tmp_path):
    path = reproduction / "report.json"
    report = json.loads(path.read_text())
    next(r for r in report["results"] if r["id"] == "figure2")["status"] = "passed"
    path.write_text(json.dumps(report))
    with pytest.raises(ValueError, match="missing source"):
        builder.build(reproduction, tmp_path / "sources.xlsx")


def test_checker_detects_modified_numeric_cell(reproduction, tmp_path):
    output = tmp_path / "sources.xlsx"
    builder.build(reproduction, output)
    metadata = json.loads(output.with_suffix(".validation.json").read_text())
    table = next(t for t in metadata["tables"] if t["csv"] == "figure2c_experimental_participants.csv")
    workbook = load_workbook(output)
    workbook[table["sheet"]].cell(table["firstDataRow"], 3).value = 99
    workbook.save(output)
    workbook.close()
    with pytest.raises(ValueError, match="Value differs"):
        builder.check(reproduction, output)


def test_checker_detects_stale_coverage_label(reproduction, tmp_path):
    output = tmp_path / "sources.xlsx"
    builder.build(reproduction, output)
    metadata = json.loads(output.with_suffix(".validation.json").read_text())
    row = next(r["row"] for r in metadata["guide"] if r["values"][0] == "Figure 2C model")
    workbook = load_workbook(output)
    workbook["Source guide"].cell(row, 2).value = "available"
    workbook.save(output)
    workbook.close()
    with pytest.raises(ValueError, match="Source guide differs"):
        builder.check(reproduction, output)


def test_figure2c_status_selects_required_tables_when_figure2b_is_blocked(reproduction, tmp_path):
    path = reproduction / "report.json"
    report = json.loads(path.read_text())
    record = next(r for r in report["results"] if r["id"] == "figure2")
    record.update(status="blocked", details={"panels": {"2B": "blocked", "2C": "passed"}})
    path.write_text(json.dumps(report))
    with pytest.raises(ValueError, match="missing source: tables/figure2c_direction_correlations"):
        builder.build(reproduction, tmp_path / "sources.xlsx")
