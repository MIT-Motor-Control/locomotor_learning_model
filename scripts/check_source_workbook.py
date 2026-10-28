"""Verify a source workbook against the reproduction report, CSV tables and checksums.

Usage: python scripts/check_source_workbook.py REPRODUCTION_DIR OUTPUT.xlsx
"""

from __future__ import annotations

import csv
import hashlib
import json
import math
import re
import sys
from pathlib import Path

from openpyxl import load_workbook


def typed(value):
    """Convert a CSV field with the same rules as ``build_source_workbook.typed``."""
    if value == "":
        return None
    if value in ("True", "False"):
        return value == "True"
    if re.fullmatch(r"[+-]?(?:\d+\.?\d*|\.\d+)(?:[Ee][+-]?\d+)?", value):
        number = float(value)
        if not math.isfinite(number):
            return value
        if re.fullmatch(r"[+-]?\d+", value) and len(value.lstrip("+-0")) > 15:
            return value
        return number
    return value


def equal(actual, expected):
    """Return whether a workbook value matches a CSV value (numbers: rtol 2e-14, atol 1e-14)."""
    if isinstance(expected, bool):
        return isinstance(actual, bool) and actual == expected
    if isinstance(expected, (float, int)):
        return (
            isinstance(actual, (float, int))
            and not isinstance(actual, bool)
            and math.isclose(actual, expected, rel_tol=2e-14, abs_tol=1e-14)
        )
    return actual == expected


def check(input_root: Path, output: Path) -> dict:
    """Verify a workbook against its sidecar metadata, report.json and the source CSVs.

    Checks the report, input-manifest and CSV checksums, every table header and cell, the
    source guide, the source index and the coverage labels of the Figure 2C model, Figure 3D
    and grid-sensitivity tables, and rejects formula or error cells.

    Args:
        input_root: Reproduction output directory containing report.json and tables/.
        output: Workbook path; the ``.validation.json`` sidecar is read from beside it.

    Returns:
        Counts of sheets, tables, verified cells and source CSVs, with ``status``.

    Raises:
        ValueError: If the metadata schema is unsupported or any check fails.
    """
    metadata = json.loads(output.with_suffix(".validation.json").read_text())
    if metadata.get("schema_version") != 3:
        raise ValueError("Unsupported workbook metadata; rebuild with build_source_workbook.py")
    input_root = input_root.resolve()
    report_payload = (input_root / "report.json").read_bytes()
    if hashlib.sha256(report_payload).hexdigest() != metadata["sourceReportSha256"]:
        raise ValueError("Reproduction report changed after workbook build")
    report = json.loads(report_payload)
    report_artifacts = {item["path"]: item["sha256"] for item in report["artifacts"]}
    if report["inputs"]["manifest_sha256"] != metadata["inputManifestSha256"]:
        raise ValueError("Workbook input-manifest provenance differs")
    workbook = load_workbook(output, read_only=False, data_only=False)
    cells = 0
    hashes = {}

    def source(name, digest):
        """Return the rows of one CSV after checking it against the sidecar and report hashes."""
        path = input_root / "tables" / name
        if not path.resolve().is_relative_to(input_root / "tables"):
            raise ValueError("Workbook metadata contains an unsafe CSV path")
        hashes[name] = hashlib.sha256(path.read_bytes()).hexdigest()
        if hashes[name] != digest or report_artifacts.get(f"tables/{name}") != digest:
            raise ValueError(f"CSV differs from the workbook or reproduction report: {name}")
        with path.open(newline="", encoding="utf-8-sig") as stream:
            return list(csv.DictReader(stream))

    for table in metadata["tables"]:
        rows = source(table["csv"], table["sourceSha256"])
        if len(rows) != table["rows"]:
            raise ValueError(f"Row count differs: {table['csv']}")
        sheet = workbook[table["sheet"]]
        headers = [
            cell.value
            for cell in next(
                sheet.iter_rows(
                    min_row=table["headerRow"],
                    max_row=table["headerRow"],
                    max_col=table["columns"],
                )
            )
        ]
        if headers != table["displayHeaders"]:
            raise ValueError(f"Headers differ: {table['sheet']}")
        actual_rows = sheet.iter_rows(
            min_row=table["firstDataRow"],
            max_row=table["firstDataRow"] + len(rows) - 1,
            max_col=table["columns"],
            values_only=True,
        )
        for i, (original, actual) in enumerate(zip(rows, actual_rows)):
            for j, field in enumerate(table["originalHeaders"]):
                expected = typed(original[field])
                if not equal(actual[j], expected):
                    raise ValueError(
                        f"Value differs: {table['sheet']} source row {i + 1}, {field}: "
                        f"{actual[j]!r} != {expected!r}"
                    )
                cells += 1
    guide = workbook["Source guide"]
    for entry in metadata["guide"]:
        actual = [guide.cell(entry["row"], c).value for c in (1, 2, 3)]
        if actual != entry["values"]:
            raise ValueError("Source guide differs from the validated coverage metadata")
    actual_index = list(
        workbook["Sources"].iter_rows(
            min_row=metadata["sourceIndexFirstDataRow"],
            max_row=metadata["sourceIndexFirstDataRow"] + len(metadata["sourceIndexRows"]) - 1,
            max_col=4,
            values_only=True,
        )
    )
    if [list(row) for row in actual_index] != metadata["sourceIndexRows"]:
        raise ValueError("Source index differs from the imported CSV records")
    # Cross-check coverage labels against the imported numerical tables.
    coverage = {row["values"][0]: row["values"][1] for row in metadata["guide"]}
    for panel, filename in (
        ("Figure 2C model", "figure2c_model_trials.csv"),
        ("Figure 3D", "figure3_d_fit_inputs.csv"),
        ("S1 Text grid sensitivity", "s1_grid_sensitivity_summary.csv"),
    ):
        label = coverage.get(panel)
        if label is None:
            raise ValueError(f"Workbook guide lacks coverage row: {panel}")
        if (label == "available") != (filename in hashes):
            raise ValueError(f"Coverage does not match imported source tables: {panel}")
    for sheet in workbook:
        for row in sheet.iter_rows():
            for cell in row:
                if cell.data_type in ("e", "f"):
                    raise ValueError(
                        f"Unexpected spreadsheet error or formula: {sheet.title}!{cell.coordinate}"
                    )
    result = {
        "sheets": len(workbook.sheetnames),
        "tables": len(metadata["tables"]),
        "verified_cells": cells,
        "source_csvs": len(hashes),
        "status": "passed",
    }
    workbook.close()
    return result


if __name__ == "__main__":
    if len(sys.argv) != 3:
        raise SystemExit(
            "Usage: python scripts/check_source_workbook.py REPRODUCTION_DIR OUTPUT.xlsx"
        )
    print(json.dumps(check(Path(sys.argv[1]), Path(sys.argv[2])), indent=2))
