"""Format numerical publication CSVs as an editor source workbook.

Usage: python scripts/build_source_workbook.py REPRODUCTION_DIR OUTPUT.xlsx
Requires only openpyxl and the Python standard library. Scientific calculations
remain in the figure pipeline; this command preserves its recorded CSV values.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import io
import json
import math
import os
import re
import tempfile
import textwrap
from pathlib import Path

from check_source_workbook import check
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

NAVY, GRAY = "253D5B", "56616E"
NUMBER = re.compile(r"[+-]?(?:\d+\.?\d*|\.\d+)(?:[Ee][+-]?\d+)?\Z")
# Each CSV remains one flat table. Related small tables share a worksheet.
SECTIONS = (
    (
        "Fig1 summary",
        "Figure 1 group summaries",
        "For 1C-D, energy change is percent and SLA change is percentage points. Figure 1E box statistics are in W/kg.",
        ("figure1_cd_summary", "figure1_e_box_summary"),
    ),
    (
        "Fig1 energy",
        "Figure 1C participant energy",
        "Short: seconds 181-360. Long and tied baseline: final 180 seconds.",
        ("figure1_energy_participants",),
    ),
    (
        "Fig1 SLA",
        "Figure 1D participant step-length asymmetry",
        "The source script's fractional-stride short window is retained. Changes are percentage points.",
        ("figure1_sla_participants",),
    ),
    (
        "Fig1 Sanchez",
        "Figure 1C-D Sanchez comparison",
        "15 participants. Short windows use shared energy-derived indices; each long window uses its final 100 samples.",
        ("figure1_sanchez_participants",),
    ),
    (
        "Fig1 timepoints",
        "Figure 1E metabolic time course",
        "Three-minute means. Adjacent paired t-tests are two-sided and uncorrected, matching the source analysis.",
        ("figure1_energy_timeseries", "figure1_paired_tests"),
    ),
    (
        "Optimization",
        "Reported trajectory-optimization values",
        "The five plotted constants are retained. The separate four-condition source check does not regenerate that series.",
        (
            "figure1_optimization",
            "s3a_trajectory_optimization",
            "figure1_optimization_source_subset",
        ),
    ),
    (
        "Fig3 counts",
        "Figure 3C fall probabilities",
        "50 trials per cell. Probability and percentage columns have distinct scales.",
        ("figure3_c_counts",),
    ),
    (
        "Fig3 trials",
        "Figure 3C simulation trials",
        "One row per simulation. A blank fall stride denotes a nonfalling trial.",
        ("figure3_trials",),
    ),
    (
        "S1 LOSO",
        "S1 held-out prediction errors",
        "Regenerated trial-average grid. The plotted p-value uses MATLAB signrank; SciPy exact is an alternative calculation.",
        ("s1_loso_persubject", "s1_loso_stats", "s1_loso_summary", "s1_loso_matlab_test"),
    ),
    (
        "S1 sensitivity",
        "S1 shared-grid sensitivity",
        "MATLAB rng(0), 200 subsamples per fold, thresholds 26/17/8. RMSE is in percentage points.",
        ("s1_grid_sensitivity_summary", "s1_grid_finite_counts", "s1_grid_sensitivity_persubject"),
    ),
    (
        "S3 model",
        "S3B model energy relative to tied walking",
        "Per-participant median over nonfalling trials, then condition mean. Shared cells are reused across participants.",
        ("s3b_model_summary", "s3b_model_participants"),
    ),
    (
        "S3 trials",
        "S3B unique model trials",
        "Tied baseline: strides 31-70. Adapted: final 100 strides. Falling trials are excluded from participant medians.",
        ("s3b_model_trials",),
    ),
    (
        "S3 experiment",
        "S3C experimental energy relative to tied walking",
        "Percent change between final 180-second means. Correlations compare trends across the three conditions.",
        ("s3c_group_means", "s3c_experimental", "s3_direction_correlations"),
    ),
)
BASE_TABLES = {
    "figure1": {
        "figure1_cd_summary",
        "figure1_e_box_summary",
        "figure1_energy_participants",
        "figure1_sla_participants",
        "figure1_energy_timeseries",
        "figure1_paired_tests",
        "figure1_optimization",
    },
    "figure3": {"figure3_c_counts", "figure3_trials"},
    "s1": {"s1_loso_persubject", "s1_loso_stats", "s1_loso_summary", "s1_loso_matlab_test"},
    "s3": {"s3a_trajectory_optimization", "s3c_experimental", "s3c_group_means"},
}


def typed(value: str):
    if value == "":
        return None
    if value in ("True", "False"):
        return value == "True"
    if NUMBER.fullmatch(value):
        number = float(value)
        if not (-float("inf") < number < float("inf")):
            return value
        # Excel has only 15 significant digits for integer identifiers.
        if re.fullmatch(r"[+-]?\d+", value) and len(value.lstrip("+-0")) > 15:
            return value
        return number
    return value


def digest(payload: bytes) -> str:
    return hashlib.sha256(payload).hexdigest()


def load_sources(root: Path) -> tuple[dict, dict, dict]:
    report_bytes = (root / "report.json").read_bytes()
    report = json.loads(report_bytes)
    if report.get("schema_version") != 1 or report.get("inputs", {}).get("errors"):
        raise ValueError("A valid reproduction report with verified inputs is required")
    records = {r["id"]: r for r in report["results"]}
    if not BASE_TABLES.keys() <= records.keys():
        raise ValueError(
            "Reproduce figure1, figure3, s1 and s3 before building the source workbook"
        )
    artifacts = {r["path"]: r["sha256"] for r in report["artifacts"]}
    wanted = {name for *_, names in SECTIONS for name in names}
    required = set().union(
        *(names for figure, names in BASE_TABLES.items() if records[figure]["status"] != "blocked")
    )
    if records["figure1"]["status"] == "passed":
        required.add("figure1_sanchez_participants")
    if records["s3"]["status"] == "passed":
        required.update(
            {
                "s3b_model_summary",
                "s3b_model_participants",
                "s3b_model_trials",
                "s3_direction_correlations",
            }
        )
    if records["s1"].get("details", {}).get("grid_sensitivity_available"):
        required.update(
            {
                "s1_grid_sensitivity_summary",
                "s1_grid_finite_counts",
                "s1_grid_sensitivity_persubject",
            }
        )
    sources = {}
    for name in sorted(wanted):
        relative = f"tables/{name}.csv"
        path = root / relative
        if not path.exists():
            if name in required:
                raise ValueError(f"Reproduction report requires missing source: {relative}")
            continue
        if not path.resolve().is_relative_to(root.resolve()):
            raise ValueError(f"Source CSV escapes reproduction folder: {relative}")
        payload = path.read_bytes()
        if artifacts.get(relative) != digest(payload):
            raise ValueError(f"Source CSV does not match its reproduction report: {relative}")
        rows = list(csv.reader(io.StringIO(payload.decode("utf-8-sig")), strict=True))
        if (
            len(rows) < 2
            or not rows[0]
            or len(set(rows[0])) != len(rows[0])
            or any(len(row) != len(rows[0]) for row in rows)
        ):
            raise ValueError(f"Empty, duplicate-column or malformed CSV: {relative}")
        sources[name] = {
            "name": f"{name}.csv",
            "headers": rows[0],
            "rows": [[typed(x) for x in row] for row in rows[1:]],
            "sha256": digest(payload),
        }
    if not sources:
        raise ValueError("No numerical source tables are available for the requested panels")
    metadata = {
        "schema_version": 2,
        "sourceReportSha256": digest(report_bytes),
        "sourceReportCreatedAt": report["created_at"],
        "inputManifestSha256": report["inputs"]["manifest_sha256"],
        "tables": [],
        "guide": [],
    }
    return records, sources, metadata


def header(name: str) -> str:
    special = {
        "delta_mps": "Speed difference (m/s)",
        "energy_change_pct": "Energy change (%)",
        "sla_change_percentage_points": "SLA change (percentage points)",
        "fall_probability": "Fall probability (0-1)",
        "fall_probability_pct": "Fall probability (%)",
        "cond_specific": "Condition-specific RMSE (percentage points)",
        "common_full": "Common full-pool RMSE (percentage points)",
        "common_size_matched": "Common size-matched RMSE (percentage points)",
        "mean_heldout_RMSE": "Mean held-out RMSE (percentage points)",
        "median_heldout_RMSE": "Median held-out RMSE (percentage points)",
        "median_condition_specific": "Median condition-specific RMSE (percentage points)",
        "median_common_size_matched": "Median common size-matched RMSE (percentage points)",
        "wilcoxon_two_sided_exact_scipy": "Alternative exact p-value (SciPy)",
        "wilcoxon_two_sided_matlab_signrank_default": "Plotted/reference p-value (MATLAB signrank)",
        "tied_energy": "Tied energy (W/kg)",
        "split_energy": "Split energy (W/kg)",
        "tied_last180_mean": "Tied final 180-second mean (W/kg)",
        "split_last180_mean": "Split final 180-second mean (W/kg)",
        "tied_stride31_70_mean": "Tied mean, strides 31-70 (model energy proxy)",
        "adapted_last100_mean": "Adapted mean, last 100 strides (model energy proxy)",
        "sample_sd": "Sample SD (percentage points)",
        "time_min": "Time (min)",
        "net_cost_j_kg_m": "Net cost (J/kg/m)",
        "stored_short_index_not_used": "Stored short index (not used)",
    }
    return special.get(
        name,
        name.replace("_w_kg", " (W/kg)")
        .replace("_mps", " (m/s)")
        .replace("_pct", " (%)")
        .replace("_", " "),
    )


def text(cell, value):
    cell.value = value
    if isinstance(value, str):
        cell.data_type = "s"  # CSV text is never executed as an Excel formula.


def heading(sheet, title: str, note: str):
    sheet.sheet_view.showGridLines = False
    text(sheet["A2"], title)
    sheet["A2"].font = Font(name="Arial", size=14, bold=True, color=NAVY)
    sheet.row_dimensions[2].height = 26
    for index, line in enumerate(textwrap.wrap(note, width=110), 4):
        text(sheet.cell(index, 1), line)
        sheet.cell(index, 1).font = Font(name="Arial", size=10, color=GRAY)
        sheet.row_dimensions[index].height = 17
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0


def add_table(sheet, source: dict, row: int, index: int, metadata: dict) -> int:
    text(
        sheet.cell(row, 1),
        f"Source: tables/{source['name']}"
        if source["sha256"]
        else "Workbook source index (metadata)",
    )
    sheet.cell(row, 1).font = Font(name="Arial", size=10, italic=True, color=GRAY)
    row += 2
    labels = [header(key) for key in source["headers"]]
    for col, label in enumerate(labels, 1):
        cell = sheet.cell(row, col)
        text(cell, label)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(right=Side(style="thin", color="FFFFFF"))
    sheet.row_dimensions[row].height = 45
    last_row = row + len(source["rows"])
    for col, key in enumerate(source["headers"], 1):
        longest = max(
            (
                len(str(values[col - 1]))
                for values in source["rows"]
                if isinstance(values[col - 1], str)
            ),
            default=0,
        )
        width = max(17, min(29, len(labels[col - 1]) * 0.55), min(100, longest * 0.95))
        letter = get_column_letter(col)
        sheet.column_dimensions[letter].width = max(
            sheet.column_dimensions[letter].width or 0, width
        )
    for number, values in enumerate(source["rows"], row + 1):
        sheet.row_dimensions[number].height = 19
        for col, value in enumerate(values, 1):
            cell = sheet.cell(number, col)
            text(cell, value)
            cell.font = Font(name="Arial", size=10, color="202B37")
            cell.alignment = Alignment(
                horizontal="right"
                if isinstance(value, (float, int)) and not isinstance(value, bool)
                else "left",
                vertical="center",
                wrap_text=isinstance(value, str),
            )
            if isinstance(value, str):
                width = sheet.column_dimensions[get_column_letter(col)].width
                lines = max(1, math.ceil(len(value) / (width * 0.9)))
                sheet.row_dimensions[number].height = max(
                    sheet.row_dimensions[number].height, 15 * lines + 4
                )
            if isinstance(value, float):
                key = source["headers"][col - 1]
                cell.number_format = (
                    "0"
                    if value.is_integer()
                    else "0.000E+00"
                    if key.startswith("p_") or "wilcoxon" in key
                    else "0.0000"
                )
    ref = f"A{row}:{get_column_letter(len(labels))}{last_row}"
    table = Table(displayName=f"SourceTable{index}", ref=ref)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2", showRowStripes=True, showColumnStripes=False
    )
    sheet.add_table(table)
    if sheet.freeze_panes is None:
        sheet.freeze_panes = f"C{row + 1}"
        sheet.print_title_rows = f"{row}:{row}"
    metadata["tables"].append(
        {
            "sheet": sheet.title,
            "csv": source["name"],
            "sourceSha256": source["sha256"],
            "headerRow": row,
            "firstDataRow": row + 1,
            "rows": len(source["rows"]),
            "columns": len(labels),
            "originalHeaders": source["headers"],
            "displayHeaders": labels,
        }
    )
    return last_row + 4


def build(input_root: Path, output: Path, *, overwrite: bool = False) -> dict:
    input_root, output = Path(input_root).resolve(), Path(output).absolute()
    if output.suffix.lower() != ".xlsx":
        raise ValueError("Output must have an .xlsx extension")
    sidecar = output.with_suffix(".validation.json")
    if output.is_symlink() or sidecar.is_symlink():
        raise ValueError("Workbook output paths must not be symlinks")
    if not overwrite and (output.exists() or sidecar.exists()):
        raise FileExistsError(
            "Workbook or validation file already exists; choose a new path or use --overwrite"
        )
    records, sources, metadata = load_sources(input_root)
    workbook = Workbook()
    guide = workbook.active
    guide.title = "Source guide"
    heading(
        guide,
        "Numerical source data",
        "PBIOLOGY-D-26-00535R2. CSVs remain the numerical sources; this workbook formats their recorded values.",
    )
    guide.sheet_properties.tabColor = NAVY
    guide_rows = [
        [
            "1C-D",
            records["figure1"]["status"],
            "Participant measurements, Sanchez comparison when available, and five reported optimization values.",
        ],
        [
            "1E",
            "available" if "figure1_energy_timeseries" in sources else "unavailable",
            "Metabolic timepoints and adjacent paired tests. Sanchez data are not part of this panel.",
        ],
        [
            "3C",
            "available"
            if "figure3_c_counts" in sources and "figure3_trials" in sources
            else "unavailable",
            "Per-cell fall counts and individual simulation trials.",
        ],
        [
            "S1",
            records["s1"]["status"],
            "Held-out errors and both labeled test methods. MATLAB signrank supplies the plotted p-value.",
        ],
        [
            "S1 sensitivity",
            "available" if "s1_grid_sensitivity_summary" in sources else "not supplied",
            "Shared-grid thresholds 26/17/8, finite-cell counts and participant folds.",
        ],
        [
            "S3A",
            "reported values" if "s3a_trajectory_optimization" in sources else "unavailable",
            "Five plotted constants. The separate source subset is not a full OpenSim regeneration.",
        ],
        [
            "S3B",
            "available" if "s3b_model_trials" in sources else "unavailable",
            "Regenerated trials at original fitted cells, participant medians and condition means.",
        ],
        [
            "S3C",
            "available" if "s3c_experimental" in sources else "unavailable",
            "Participant energy changes relative to the tied baseline.",
        ],
    ]
    for row, values in enumerate([["Panel", "Coverage", "Contents and limits"], *guide_rows], 7):
        for col, value in enumerate(values, 1):
            cell = guide.cell(row, col)
            text(cell, value)
            cell.font = Font(
                name="Arial", size=10, bold=row == 7, color="FFFFFF" if row == 7 else "202B37"
            )
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            if row == 7:
                cell.fill = PatternFill("solid", fgColor=NAVY)
        guide.row_dimensions[row].height = 38 if row != 7 else 25
    for letter, width in zip("ABC", (23, 23, 100)):
        guide.column_dimensions[letter].width = width
    notes = [
        f"Reproduction: {metadata['sourceReportCreatedAt']}",
        "Percent-change values already use percent units: a stored value of 20 means 20%.",
        "Blank cells remain blank. Literal nan entries identify unavailable values, not zeros.",
        "Participant indices are condition-local unless an explicit participant identifier is present.",
        "No measurements, fits or statistical tests are recalculated by this workbook.",
    ]
    for row, note in enumerate(notes, 18):
        text(guide.cell(row, 1), note)
        guide.cell(row, 1).font = Font(name="Arial", size=10, color=GRAY)
        guide.row_dimensions[row].height = 18
    metadata["guide"] = [{"row": row, "values": values} for row, values in enumerate(guide_rows, 8)]
    for name, title, note, names in SECTIONS:
        available = [sources[key] for key in names if key in sources]
        if not available:
            continue
        sheet = workbook.create_sheet(name)
        heading(sheet, title, note)
        row = 8
        for source in available:
            row = add_table(sheet, source, row, len(metadata["tables"]) + 1, metadata)
    index = workbook.create_sheet("Sources")
    heading(
        index,
        "Source index",
        "Each table is copied from its listed CSV. SHA-256 values identify the exact numerical inputs.",
    )
    index_source = {
        "name": "source index (metadata)",
        "sha256": "",
        "headers": ["worksheet", "source_csv", "records", "source_sha256"],
        "rows": [[t["sheet"], t["csv"], t["rows"], t["sourceSha256"]] for t in metadata["tables"]],
    }
    count = len(metadata["tables"])
    add_table(index, index_source, 8, count + 1, metadata)
    metadata["tables"].pop()  # The index is metadata, not another scientific CSV.
    metadata["sourceIndexRows"] = index_source["rows"]
    metadata["sourceIndexFirstDataRow"] = 11
    text(
        index.cell(len(index_source["rows"]) + 14, 1),
        f"Reproduction report SHA-256: {metadata['sourceReportSha256']}",
    )
    text(
        index.cell(len(index_source["rows"]) + 15, 1),
        f"Input manifest SHA-256: {metadata['inputManifestSha256']}",
    )
    output.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="source-workbook-", dir=output.parent) as scratch:
        staged = Path(scratch) / output.name
        workbook.save(staged)
        workbook.close()
        staged.with_suffix(".validation.json").write_text(json.dumps(metadata, indent=2) + "\n")
        result = check(input_root, staged)
        os.replace(staged, output)
        os.replace(staged.with_suffix(".validation.json"), sidecar)
    return {"workbook": str(output), **result}


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("reproduction_dir", type=Path)
    parser.add_argument("output", type=Path)
    parser.add_argument("--overwrite", action="store_true")
    args = parser.parse_args(argv)
    try:
        result = build(args.reproduction_dir, args.output, overwrite=args.overwrite)
    except (OSError, ValueError, KeyError, AssertionError) as error:
        parser.exit(1, f"Cannot build source workbook: {error}\n")
    print(json.dumps(result, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
